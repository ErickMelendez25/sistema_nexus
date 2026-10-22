import asyncio
from playwright.async_api import async_playwright
from bs4 import BeautifulSoup
import os
from openpyxl import Workbook, load_workbook

# Archivo final en formato Excel puro (.xlsx)
ARCHIVO_EXCEL = "convocatorias_seace.xlsx"

def guardar_en_excel_real(datos_fila):
    try:
        if not os.path.exists(ARCHIVO_EXCEL):
            wb = Workbook()
            ws = wb.active
            ws.title = "Convocatorias SEACE"
            ws.append([
                "N°", "Nombre o Sigla de Entidad", "Tipo de Selección", 
                "Nro. Selección", "Descripción del Objeto", "Año", 
                "Monto / Valor", "Estado / Fecha", "Datos Adicionales"
            ])
            wb.save(ARCHIVO_EXCEL)
        
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb.active
        ws.append(datos_fila)
        wb.save(ARCHIVO_EXCEL)
    except Exception as e:
        print(f"❌ Error al guardar en Excel: {e}")

def limpiar_y_mostrar_tabla(xml_texto):
    try:
        soup_xml = BeautifulSoup(xml_texto, "xml")
        actualizaciones = soup_xml.find_all("update")
        
        for update in actualizaciones:
            update_id = update.get("id", "")
            
            if "pnlGrdResultadosProcesos" in update_id or "dtProcesos" in update_id:
                html_interno = update.text
                soup_html = BeautifulSoup(html_interno, "html.parser")
                filas = soup_html.find_all("tr")
                
                if len(filas) <= 1:
                    continue
                    
                print("\n========================================================")
                print(f"📊 [PROCESANDO NUEVOS DATOS] Se encontraron {len(filas)} filas.")
                print("========================================================")
                
                contador_real = 1
                for idx, fila in enumerate(filas):
                    columnas = [td.get_text(strip=True) for td in fila.find_all("td")]
                    if columnas and len(columnas) > 3:
                        texto_linea = " | ".join(columnas)
                        if "Código Unico" in texto_linea or "No se encontraron" in texto_linea:
                            continue
                            
                        print(f"🔹 Convocatoria #{contador_real}: {texto_linea[:120]}...")
                        guardar_en_excel_real(columnas)
                        contador_real += 1
                        
    except Exception as e:
        print(f"❌ Error al procesar el texto: {e}")

async def interceptar_seace(response):
    if "buscadorPublico.xhtml" in response.url:
        try:
            texto_respuesta = await response.text()
            if "pnlGrdResultadosProcesos" in texto_respuesta or "dtProcesos" in texto_respuesta:
                limpiar_y_mostrar_tabla(texto_respuesta)
        except Exception:
            pass

async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False)
        context = await browser.new_context()
        page = await context.new_page()
        
        page.on("response", interceptar_seace)
        
        print("🔗 Navegando al buscador público del SEACE...")
        await page.goto("https://seace.gob.pe")
        
        print("\n💡 INSTRUCCIONES:")
        print("1. Llena los filtros y dale 'Buscar'.")
        print("2. Presiona ENTER aquí en la consola para iniciar la ráfaga de velocidad extrema.")
        
        input("\n👉 Presiona [ENTER] aquí cuando ya veas los primeros resultados...")
        
        print("\n🚀 [VELOCIDAD EXTREMA ACTIVADA] Extrayendo por inyección de JavaScript...")
        
        pag_actual = 1
        while True:
            # Verificamos si el botón 'Siguiente' en el HTML está deshabilitado
            esta_deshabilitado = await page.eval_on_selector(
                ".ui-paginator-next", 
                "el => el.classList.contains('ui-state-disabled')"
            )
            
            if esta_deshabilitado:
                print("\n🏁 [FIN DEL PROCESO] Se completaron todas las páginas disponibles.")
                break
                
            print(f"⚡ Disparando ráfaga asíncrona hacia la siguiente página (Desde Pág. {pag_actual})...")
            
            # EJECUCIÓN DIRECTA POR JAVASCRIPT: Forzamos el evento AJAX de PrimeFaces del SEACE.
            # Esto simula el envío del formulario instantáneamente sin mover el cursor ni hacer clics visuales.
            await page.evaluate("""
                const btn = document.querySelector('.ui-paginator-next');
                if (btn) { btn.click(); }
            """)
            
            # Reducimos drásticamente el tiempo de espera plano a solo 1.2 segundos (1200ms)
            # que es el tiempo promedio de respuesta de la red del gobierno.
            await page.wait_for_timeout(1200)
            pag_actual += 1
                
        print(f"\n🎉 ¡Extracción masiva finalizada a toda velocidad! Revisa '{ARCHIVO_EXCEL}'")

if __name__ == "__main__":
    asyncio.run(main())
