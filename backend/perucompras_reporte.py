"""
Helbot - perucompras_reporte.py
Reporte de cobertura (SOLO LECTURA) en modo rápido: varios navegadores en paralelo,
sin esperas fijas, sin imágenes/fuentes y con atajo por región.
NO modifica plazos ni presiona "Enviar". Genera un Excel profesional.
"""
import io
import os
import re
import queue
import logging
import threading
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor

from perucompras_pw import (
    PWTimeout, abrir_pagina, reservar_usuario, esperar_ajax, opciones_select, sesion_caida, captura,
    ir_a, select2_seleccionar,
)
from perucompras_plazo import _obtener_arbol, _preparar, URL_PLAZO

logger = logging.getLogger("helbot.perucompras_reporte")

CON = "con_resultados"
SIN = "sin_resultados"
ERR = "error"

# Navegadores en paralelo. Cada Chromium consume ~300 MB de RAM en tu VPS.
# Se puede cambiar sin tocar código con la variable de entorno REPORTE_WORKERS.
WORKERS_DEFECTO = int(os.getenv("REPORTE_WORKERS", "4"))
# Atajo por región (busca "Todas las provincias" primero). 0 = desactivarlo.
ATAJO_REGION = os.getenv("REPORTE_ATAJO_REGION", "1") != "0"
RE_TODAS = re.compile(r"^\s*(-+\s*)?(todas?|todos?)\b", re.I)

_JS_LIMPIAR = r"""() => {
  const t = document.getElementById('tbData'); if (t) t.innerHTML = '';
  const i = document.getElementById('divNumRegistrosPorPagina'); if (i) i.innerText = '';
}"""

_JS_LISTO = r"""() => {
  const tb = document.getElementById('tbData');
  const txt = (tb && tb.innerText) || '';
  return /No se encontraron fichas/i.test(txt) || !!document.querySelector('#tbData input.PlazoEntrega');
}"""

_JS_INFO_LISTA = r"""() => /de\s+\d+/.test(((document.getElementById('divNumRegistrosPorPagina')||{}).innerText)||'')"""

_JS_ESTADO = r"""() => {
  const tb = document.getElementById('tbData');
  const txt = (tb && tb.innerText) || '';
  return {
    sin: /No se encontraron fichas/i.test(txt),
    inputs: document.querySelectorAll('#tbData input.PlazoEntrega').length,
    filas: document.querySelectorAll('#tbData tr.FilaDatos').length,
    info: ((document.getElementById('divNumRegistrosPorPagina')||{}).innerText) || ''
  };
}"""


def _acelerar(page):
    """Corta descargas que no aportan nada a la consulta."""
    try:
        page.route(
            "**/*",
            lambda route: route.abort()
            if route.request.resource_type in ("image", "font", "media")
            else route.continue_(),
        )
    except Exception as e:
        logger.warning(f"No se pudo activar el bloqueo de recursos: {e}")


_JS_FIRMA = r"""id => [...document.querySelectorAll('#'+id+' option')]
  .filter(o => o.value && o.value !== '0').map(o => o.value).join('|')"""

_JS_CAMBIO = r"""([id, antes]) => {
  const f = [...document.querySelectorAll('#'+id+' option')]
    .filter(o => o.value && o.value !== '0').map(o => o.value).join('|');
  return f !== '' && f !== antes;
}"""


def _firma(page, sel_id: str) -> str:
    return page.evaluate(_JS_FIRMA, sel_id)


def _esperar_cambio(page, sel_id: str, firma_antes: str, timeout_ms: int) -> bool:
    """Espera a que la lista de un <select> cambie. Reemplaza las esperas fijas."""
    try:
        page.wait_for_function(_JS_CAMBIO, arg=[sel_id, firma_antes], timeout=timeout_ms)
        return True
    except PWTimeout:
        return False


def _preparar_rapido(page, acuerdo_id, catalogo_id, categoria_id):
    ir_a(page, URL_PLAZO)
    page.wait_for_selector("#cboAcuerdo", state="attached")

    antes = _firma(page, "cboCatalogo")
    select2_seleccionar(page, "cboAcuerdo", acuerdo_id)
    if not _esperar_cambio(page, "cboCatalogo", antes, 8000):
        esperar_ajax(page)

    antes = _firma(page, "cboCategoria")
    page.select_option("#cboCatalogo", str(catalogo_id))
    if not _esperar_cambio(page, "cboCategoria", antes, 8000):
        esperar_ajax(page)

    antes = _firma(page, "cboRegion")
    page.select_option("#cboCategoria", str(categoria_id))
    if not _esperar_cambio(page, "cboRegion", antes, 6000):
        esperar_ajax(page)


def _elegir_region(page, valor: str):
    antes = _firma(page, "cboProvincia")
    page.select_option("#cboRegion", valor)
    if not _esperar_cambio(page, "cboProvincia", antes, 6000):
        esperar_ajax(page, 150)


_MUESTRAS = [0]


def _log_muestra(resp):
    """Guarda en el log 3 peticiones reales (para armar el modo directo, sin navegador)."""
    if _MUESTRAS[0] >= 3:
        return
    _MUESTRAS[0] += 1
    try:
        req = resp.request
        logger.warning(
            "TURBO-MUESTRA url=%s metodo=%s content-type=%s post=%s respuesta[:500]=%r",
            resp.url, req.method, req.headers.get("content-type"),
            (req.post_data or "")[:800], resp.text()[:500],
        )
    except Exception as e:
        logger.warning(f"TURBO-MUESTRA no se pudo leer: {e}")


TURBO = os.getenv("REPORTE_TURBO", "1") != "0"   # REPORTE_TURBO=0 lo apaga
TOKEN_P = "@@P@@"
_TL = threading.local()
_HDR_OK = ("content-type", "x-requested-with", "accept")

_JS_LOTE = r"""async ([url, headers, bodies]) => {
  const out = new Array(bodies.length);
  let i = 0;
  async function w() {
    while (true) {
      const k = i++;
      if (k >= bodies.length) return;
      try {
        const r = await fetch(url, {method: 'POST', headers: headers, body: bodies[k], credentials: 'include'});
        out[k] = {s: r.status, t: await r.text()};
      } catch (e) {
        out[k] = {s: 0, t: ''};
      }
    }
  }
  await Promise.all(Array.from({length: Math.min(5, bodies.length)}, w));
  return out;
}"""


class _Turbo:
    def __init__(self):
        self.activo = TURBO
        self.plantilla = None
        self.url = None
        self.headers = None
        self.firma_vacia = None
        self.validado = False


def _apagar(turbo, motivo: str):
    turbo.activo = False
    logger.warning(f"TURBO desactivado para esta categoría: {motivo}")


def _guardar_peticion(resp):
    try:
        req = resp.request
        cab = {k: v for k, v in req.all_headers().items()
               if k.lower() in _HDR_OK or "token" in k.lower() or "verification" in k.lower()}
        _TL.req = {"url": resp.url, "post": req.post_data or "", "headers": cab}
    except Exception:
        _TL.req = None


def _contar_directo(texto: str):
    if "Nro^Imagen" not in texto:
        return None
    partes = texto.split("¯")
    if len(partes) < 2:
        return None
    n = len([r for r in partes[1].split("¬") if r.strip()])
    return n if n > 0 else None


def _directo_lote(page, turbo, valores: list) -> list:
    bodies = [turbo.plantilla.replace(TOKEN_P, v) for v in valores]
    return page.evaluate(_JS_LOTE, [turbo.url, turbo.headers, bodies])


def _clasificar(turbo, item):
    if item["s"] != 200:
        return None
    t = item["t"]
    n = _contar_directo(t)
    if n:
        return CON, n
    if turbo.firma_vacia is not None and t.strip() == turbo.firma_vacia:
        return SIN, 0
    return None


def _turbo_aprender(page, turbo, pv_value: str, estado: str, n: int):
    """Tras una consulta real por navegador: aprende la plantilla y la valida contra el navegador."""
    try:
        if turbo.plantilla is None:
            r = getattr(_TL, "req", None)
            if not r or not r["post"] or r["post"].count(pv_value) != 1:
                _apagar(turbo, "no se pudo identificar el código de provincia en la petición")
                return
            turbo.plantilla = r["post"].replace(pv_value, TOKEN_P)
            turbo.url, turbo.headers = r["url"], r["headers"]
        if turbo.validado and turbo.firma_vacia is not None:
            return
        item = _directo_lote(page, turbo, [pv_value])[0]
        if item["s"] != 200:
            _apagar(turbo, f"la petición directa respondió {item['s']}")
            return
        if estado == CON and not turbo.validado:
            if _contar_directo(item["t"]) == n:
                turbo.validado = True
            else:
                _apagar(turbo, f"conteo directo distinto al del navegador ({_contar_directo(item['t'])} vs {n})")
        elif estado == SIN and turbo.firma_vacia is None:
            t = item["t"].strip()
            if len(t) < 600 and _contar_directo(t) is None:
                turbo.firma_vacia = t
    except Exception as e:
        _apagar(turbo, f"error aprendiendo la plantilla: {str(e)[:80]}")


def _lote_turbo(page, turbo, ac, ct, ca, reg, pendientes, registrar):
    """Resuelve en directo las provincias de una región. Devuelve las que quedaron sin resolver
    (van por navegador) o None si la verificación contra el navegador falló."""
    try:
        items = _directo_lote(page, turbo, [p["value"] for p in pendientes])
        clas = [_clasificar(turbo, it) for it in items]
        idx = next((i for i in range(len(clas) - 1, -1, -1) if clas[i] is not None), None)
        if idx is None:
            return list(pendientes)
        page.select_option("#cboProvincia", pendientes[idx]["value"])
        est, n, _m = _consultar(page)
        if (est, n) != (clas[idx][0], clas[idx][1]):
            _apagar(turbo, f"verificación fallida en {reg['texto']}: directo={clas[idx]} navegador={(est, n)}")
            return None
    except Exception as e:
        _apagar(turbo, f"error en el lote: {str(e)[:80]}")
        return None
    restantes = []
    for p, c in zip(pendientes, clas):
        if c is None:
            restantes.append(p)
            continue
        est, n = c
        txt = f"{n} fichas encontradas" if est == CON else "Sin fichas"
        registrar(ac, ct, ca, reg["texto"], p["texto"], est, n, txt + " (modo directo)")
    return restantes

def _consultar(page) -> tuple[str, int, str]:
    """Hace la búsqueda y devuelve (estado, total_fichas, mensaje). No modifica nada."""
    page.evaluate(_JS_LIMPIAR)
    try:
        with page.expect_response(lambda r: "consultaMejoraPlazoEntrega" in r.url, timeout=30000) as _info_resp:
            page.get_by_text("Iniciar Búsqueda").first.click()
    except PWTimeout:
        raise RuntimeError("La búsqueda no respondió en 30s")
    _log_muestra(_info_resp.value)
    _guardar_peticion(_info_resp.value)

    try:
        page.wait_for_function(_JS_LISTO, timeout=10000)
    except PWTimeout:
        pass

    est = page.evaluate(_JS_ESTADO)
    if est["sin"]:
        return SIN, 0, "Sin fichas"
    if est["inputs"] == 0:
        raise RuntimeError("La tabla no cargó (ni fichas ni mensaje de vacío)")

    # El texto "Mostrando 1 a 10 de 49 Registros" a veces llega un instante después
    try:
        page.wait_for_function(_JS_INFO_LISTA, timeout=1500)
        est = page.evaluate(_JS_ESTADO)
    except PWTimeout:
        pass
    m = re.search(r"de\s+(\d+)", est["info"])
    total = int(m.group(1)) if m else est["filas"]
    return CON, total, f"{total} fichas encontradas"


def _procesar_categoria(page, ac, ct, ca, dep_f, prov_f, atajo, lock, progreso, registrar, abortar):
    nombre = ca["nombre"]
    turbo = _Turbo()

    def preparar():
        _preparar_rapido(page, ac["acuerdo_id"], ct["id"], ca["id"])

    try:
        preparar()
        regiones = opciones_select(page, "cboRegion")
    except Exception as e:
        if sesion_caida(page):
            raise
        captura(page, f"reporte_cat_{ca['id']}")
        with lock:
            progreso["total_combinaciones"] += 1
        registrar(ac, ct, ca, "-", "-", ERR, 0, f"No se pudo preparar la categoría: {str(e)[:100]}")
        return
    if dep_f:
        regiones = [r for r in regiones if r["value"] == dep_f]

    sucio = False
    for reg in regiones:
        if abortar.is_set():
            return
        try:
            if sucio:
                preparar()
                sucio = False
            _elegir_region(page, reg["value"])
            provs = opciones_select(page, "cboProvincia")
            if not provs:
                page.wait_for_timeout(1000)
                provs = opciones_select(page, "cboProvincia")
        except Exception as e:
            if sesion_caida(page):
                raise
            sucio = True
            with lock:
                progreso["total_combinaciones"] += 1
            registrar(ac, ct, ca, reg["texto"], "-", ERR, 0, f"Error leyendo provincias: {str(e)[:100]}")
            continue

        todas = next((p for p in provs if RE_TODAS.match(p["texto"] or "")), None)
        provs = [p for p in provs if p is not todas]
        if prov_f:
            provs = [p for p in provs if p["value"] == prov_f]
        with lock:
            progreso["total_combinaciones"] += len(provs)

        # ---- Atajo: región completa vacía => todas sus provincias vacías ----
        if atajo and todas and len(provs) > 1 and not prov_f:
            try:
                progreso["combinacion_actual"] = f"{nombre} · {reg['texto']} (región completa)"
                page.select_option("#cboProvincia", todas["value"])
                est, _n, _m = _consultar(page)
                if est == SIN:
                    for pv in provs:
                        registrar(ac, ct, ca, reg["texto"], pv["texto"], SIN, 0,
                                  "Sin fichas (región completa sin resultados)")
                    continue
            except Exception as e:
                if sesion_caida(page):
                    raise
                sucio = True  # si falla el atajo, se sigue provincia por provincia

        pendientes = list(provs)
        lote_probado = False
        while pendientes:
            if abortar.is_set():
                return
            if turbo.activo and turbo.validado and not lote_probado and len(pendientes) >= 3:
                lote_probado = True
                progreso["combinacion_actual"] = f"{nombre} · {reg['texto']} (modo directo)"
                resto = _lote_turbo(page, turbo, ac, ct, ca, reg, pendientes, registrar)
                if resto is None:
                    sucio = True
                else:
                    pendientes = resto
                continue

            pv = pendientes.pop(0)
            progreso["combinacion_actual"] = f"{nombre} · {reg['texto']} / {pv['texto']}"
            estado, n, msg = ERR, 0, "Sin intentar"
            for intento in range(3):
                try:
                    if sucio:
                        preparar()
                        _elegir_region(page, reg["value"])
                        sucio = False
                    page.select_option("#cboProvincia", pv["value"])
                    estado, n, msg = _consultar(page)
                    if turbo.activo:
                        _turbo_aprender(page, turbo, pv["value"], estado, n)
                    break
                except Exception as e:
                    if sesion_caida(page):
                        raise RuntimeError("La sesión de Perú Compras se cayó durante el reporte")
                    sucio = True
                    estado, n, msg = ERR, 0, f"Falló tras {intento + 1} intento(s): {str(e)[:110]}"
                    if intento == 2:
                        captura(page, f"reporte_{ca['id']}_{pv['value']}")
            registrar(ac, ct, ca, reg["texto"], pv["texto"], estado, n, msg)


def ejecutar_reporte_resultados(pc_session, progreso: dict, seleccion: dict | None, resultados: list) -> list[dict]:
    """Llena `resultados` EN SITIO (así, si se cae a la mitad, el reporte parcial sigue disponible)."""
    sel = seleccion or {}
    acuerdos_ok = set(sel.get("acuerdos") or [])
    cats_ok = {(int(c[0]), int(c[1])) for c in (sel.get("categorias") or [])}
    dep_f = sel.get("departamento_codigo") or None
    prov_f = sel.get("provincia_codigo") or None
    n_workers = max(1, min(6, int(sel.get("workers") or WORKERS_DEFECTO)))

    progreso.update({
        "combinacion_actual": "Abriendo navegadores...", "combinaciones_completadas": 0,
        "total_combinaciones": 0, "con_resultados": 0, "sin_resultados": 0,
        "combos_con_error": 0, "fichas_encontradas": 0,
        "categorias_total": 0, "categorias_procesadas": 0,
    })

    lock = threading.Lock()
    abortar = threading.Event()
    fatales: list[str] = []

    def registrar(ac, ct, ca, dep, prov, estado, n, msg):
        with lock:
            resultados.append({
                "acuerdo": ac["acuerdo"], "catalogo": ct["nombre"], "categoria": ca["nombre"],
                "departamento": dep, "provincia": prov, "fichas": n, "estado": estado, "mensaje": msg,
            })
            progreso["combinaciones_completadas"] += 1
            if estado == CON:
                progreso["con_resultados"] += 1
                progreso["fichas_encontradas"] += n
            elif estado == SIN:
                progreso["sin_resultados"] += 1
            else:
                progreso["combos_con_error"] += 1

    # El árbol se lee UNA vez (si ya está en caché es instantáneo); los workers lo comparten
    arbol = _obtener_arbol(pc_session)
    combos = []
    for ac in arbol:
        if acuerdos_ok and ac["acuerdo"] not in acuerdos_ok:
            continue
        for ct in ac["catalogos"]:
            for ca in ct["categorias"]:
                if cats_ok and (int(ct["id"]), int(ca["id"])) not in cats_ok:
                    continue
                combos.append((ac, ct, ca))
    if not combos:
        raise RuntimeError("Ninguna categoría seleccionada existe para este usuario.")
    progreso["categorias_total"] = len(combos)

    cola: "queue.Queue" = queue.Queue()
    for c in combos:
        cola.put(c)
    n_workers = min(n_workers, len(combos))
    atajo = ATAJO_REGION and not prov_f

    def worker(num: int):
        try:
            with abrir_pagina(pc_session, usar_lock=False) as (page, _dialogos):
                _acelerar(page)
                while not abortar.is_set():
                    try:
                        ac, ct, ca = cola.get_nowait()
                    except queue.Empty:
                        return
                    _procesar_categoria(page, ac, ct, ca, dep_f, prov_f, atajo,
                                        lock, progreso, registrar, abortar)
                    with lock:
                        progreso["categorias_procesadas"] += 1
        except Exception as e:
            logger.warning(f"Worker {num} se detuvo: {e}")
            with lock:
                fatales.append(str(e))
            abortar.set()

    with reservar_usuario(pc_session):
        with ThreadPoolExecutor(max_workers=n_workers) as pool:
            futuros = [pool.submit(worker, i + 1) for i in range(n_workers)]
            for f in futuros:
                f.result()

    if fatales:
        raise RuntimeError(fatales[0])
    if progreso["total_combinaciones"] == 0:
        raise RuntimeError("No se armó ninguna combinación categoría/región/provincia.")
    return resultados

# ======================================================================
# Excel profesional
# ======================================================================
def _situacion(g: dict) -> str:
    if g["err"] > 0 and g["con"] == 0 and g["sin"] == 0:
        return "ERROR AL CONSULTAR"
    if g["con"] == 0:
        return "SIN RESULTADOS EN TODA LA REGIÓN"
    if g["sin"] == 0 and g["err"] == 0:
        return "CON RESULTADOS EN TODAS LAS PROVINCIAS"
    return "RESULTADOS PARCIALES"


def generar_excel(resultados: list[dict], meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    AZUL = "312E81"
    f_titulo = Font(name="Calibri", size=16, bold=True, color=AZUL)
    f_sub = Font(name="Calibri", size=10, italic=True, color="64748B")
    f_head = Font(name="Calibri", bold=True, color="FFFFFF")
    f_bold = Font(name="Calibri", bold=True)
    fill_head = PatternFill("solid", fgColor=AZUL)
    fill_kpi = PatternFill("solid", fgColor="EEF2FF")
    fill_ok = PatternFill("solid", fgColor="D1FAE5")
    fill_sin = PatternFill("solid", fgColor="FEE2E2")
    fill_err = PatternFill("solid", fgColor="FEF3C7")
    lado = Side(style="thin", color="CBD5E1")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    ETQ = {CON: "CON RESULTADOS", SIN: "SIN RESULTADOS", ERR: "ERROR"}
    FILL_ESTADO = {CON: fill_ok, SIN: fill_sin, ERR: fill_err}

    filas = sorted(resultados, key=lambda r: (r["acuerdo"], r["catalogo"], r["categoria"], r["departamento"], r["provincia"]))

    def tabla(ws, fila, encabezados, datos, pct_cols=(), col_estado=None, claves_estado=None):
        for j, h in enumerate(encabezados, 1):
            c = ws.cell(row=fila, column=j, value=h)
            c.font, c.fill, c.border = f_head, fill_head, borde
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, fila_datos in enumerate(datos, fila + 1):
            for j, v in enumerate(fila_datos, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = borde
                c.alignment = Alignment(vertical="top")
                if j in pct_cols:
                    c.number_format = "0.0%"
                if col_estado and j == col_estado and claves_estado:
                    c.fill = claves_estado[i - fila - 1]
                    c.font = f_bold
        return fila + len(datos) + 1

    def auto_ancho(ws, minimo=12, maximo=55):
        anchos: dict[int, int] = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    anchos[c.column] = max(anchos.get(c.column, 0), len(str(c.value)))
        for col, w in anchos.items():
            ws.column_dimensions[get_column_letter(col)].width = max(minimo, min(maximo, w + 2))

    # ---- agrupaciones ----
    por_cat: dict[tuple, dict] = {}
    por_reg: dict[tuple, dict] = {}
    for r in filas:
        for dic, clave in ((por_cat, (r["acuerdo"], r["catalogo"], r["categoria"])),
                           (por_reg, (r["acuerdo"], r["catalogo"], r["categoria"], r["departamento"]))):
            g = dic.setdefault(clave, {"prov": 0, "con": 0, "sin": 0, "err": 0, "fichas": 0})
            g["prov"] += 1
            if r["estado"] == CON:
                g["con"] += 1
                g["fichas"] += r["fichas"]
            elif r["estado"] == SIN:
                g["sin"] += 1
            else:
                g["err"] += 1

    n_con = sum(1 for r in filas if r["estado"] == CON)
    n_sin = sum(1 for r in filas if r["estado"] == SIN)
    n_err = sum(1 for r in filas if r["estado"] == ERR)
    fichas_tot = sum(r["fichas"] for r in filas)
    cobertura = (n_con / (n_con + n_sin)) if (n_con + n_sin) else 0
    regiones_vacias = sum(1 for g in por_reg.values() if _situacion(g) == "SIN RESULTADOS EN TODA LA REGIÓN")

    wb = Workbook()

    # ================= Hoja 1: Resumen =================
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Reporte de cobertura de fichas — Perú Compras"
    ws["A1"].font = f_titulo
    ws["A2"] = "Plazo de entrega máximo · Consulta de resultados por región y provincia (solo lectura)"
    ws["A2"].font = f_sub
    meta_filas = [
        ("Usuario Perú Compras", meta.get("usuario") or "-"),
        ("Reporte generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Inicio de la consulta", (meta.get("iniciado_en") or "-")[:19].replace("T", " ")),
        ("Fin de la consulta", (meta.get("terminado_en") or "-")[:19].replace("T", " ")),
    ]
    r0 = 4
    for i, (k, v) in enumerate(meta_filas):
        ws.cell(row=r0 + i, column=1, value=k).font = f_bold
        ws.cell(row=r0 + i, column=2, value=v)

    kpis = [
        ("Acuerdos marco consultados", len({r["acuerdo"] for r in filas})),
        ("Catálogos consultados", len({(r["acuerdo"], r["catalogo"]) for r in filas})),
        ("Categorías consultadas", len(por_cat)),
        ("Combinaciones provincia consultadas", len(filas)),
        ("Provincias CON resultados", n_con),
        ("Provincias SIN resultados", n_sin),
        ("Provincias con error de consulta", n_err),
        ("Cobertura (con resultados / consultadas sin error)", cobertura),
        ("Total de fichas encontradas", fichas_tot),
        ("Regiones evaluadas (por categoría)", len(por_reg)),
        ("Regiones SIN resultados en ninguna provincia", regiones_vacias),
    ]
    r1 = r0 + len(meta_filas) + 1
    ws.cell(row=r1, column=1, value="INDICADORES GENERALES").font = Font(bold=True, size=12, color=AZUL)
    for i, (k, v) in enumerate(kpis, r1 + 1):
        a = ws.cell(row=i, column=1, value=k)
        b = ws.cell(row=i, column=2, value=v)
        a.fill = b.fill = fill_kpi
        a.border = b.border = borde
        a.font = f_bold
        b.alignment = Alignment(horizontal="right")
        if isinstance(v, float):
            b.number_format = "0.0%"

    r2 = r1 + len(kpis) + 3
    ws.cell(row=r2 - 1, column=1, value="DETALLE POR CATEGORÍA").font = Font(bold=True, size=12, color=AZUL)
    datos_cat = []
    for (ac, ct, ca), g in sorted(por_cat.items()):
        regs = [x for k, x in por_reg.items() if k[:3] == (ac, ct, ca)]
        vacias = sum(1 for x in regs if _situacion(x) == "SIN RESULTADOS EN TODA LA REGIÓN")
        cob = (g["con"] / (g["con"] + g["sin"])) if (g["con"] + g["sin"]) else 0
        datos_cat.append([ac, ct, ca, g["prov"], g["con"], g["sin"], g["err"], g["fichas"], cob, len(regs), vacias])
    tabla(ws, r2, ["Acuerdo marco", "Catálogo", "Categoría", "Provincias consultadas", "Con resultados",
                   "Sin resultados", "Con error", "Fichas encontradas", "Cobertura",
                   "Regiones evaluadas", "Regiones sin resultados"], datos_cat, pct_cols=(9,))
    auto_ancho(ws, 14, 60)
    ws.column_dimensions["A"].width = 55
    ws.sheet_view.showGridLines = False

    # ================= Hoja 2: Regiones =================
    ws2 = wb.create_sheet("Regiones")
    orden = {"SIN RESULTADOS EN TODA LA REGIÓN": 0, "ERROR AL CONSULTAR": 1,
             "RESULTADOS PARCIALES": 2, "CON RESULTADOS EN TODAS LAS PROVINCIAS": 3}
    regs_ord = sorted(por_reg.items(), key=lambda kv: (orden[_situacion(kv[1])], kv[0]))
    datos_reg, fills_reg = [], []
    for (ac, ct, ca, dep), g in regs_ord:
        sit = _situacion(g)
        datos_reg.append([ac, ct, ca, dep, g["prov"], g["con"], g["sin"], g["err"], g["fichas"], sit])
        fills_reg.append(fill_sin if orden[sit] == 0 else fill_err if orden[sit] == 1
                         else fill_ok if orden[sit] == 3 else PatternFill("solid", fgColor="FFF7ED"))
    tabla(ws2, 1, ["Acuerdo marco", "Catálogo", "Categoría", "Región / Departamento", "Provincias consultadas",
                   "Con resultados", "Sin resultados", "Con error", "Fichas", "Situación de la región"],
          datos_reg, col_estado=10, claves_estado=fills_reg)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    auto_ancho(ws2)

    # ================= Hoja 3: Sin resultados =================
    ws3 = wb.create_sheet("Sin resultados")
    datos_sin = [[r["acuerdo"], r["catalogo"], r["categoria"], r["departamento"], r["provincia"]]
                 for r in filas if r["estado"] == SIN]
    tabla(ws3, 1, ["Acuerdo marco", "Catálogo", "Categoría", "Región / Departamento", "Provincia"], datos_sin)
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = ws3.dimensions
    auto_ancho(ws3)

    # ================= Hoja 4: Con resultados =================
    ws4 = wb.create_sheet("Con resultados")
    datos_con = [[r["acuerdo"], r["catalogo"], r["categoria"], r["departamento"], r["provincia"], r["fichas"]]
                 for r in filas if r["estado"] == CON]
    tabla(ws4, 1, ["Acuerdo marco", "Catálogo", "Categoría", "Región / Departamento", "Provincia", "Fichas"], datos_con)
    ws4.freeze_panes = "A2"
    ws4.auto_filter.ref = ws4.dimensions
    auto_ancho(ws4)

    # ================= Hoja 5: Detalle completo =================
    ws5 = wb.create_sheet("Detalle completo")
    datos_det = [[r["acuerdo"], r["catalogo"], r["categoria"], r["departamento"], r["provincia"],
                  r["fichas"], ETQ[r["estado"]], r["mensaje"]] for r in filas]
    fills_det = [FILL_ESTADO[r["estado"]] for r in filas]
    tabla(ws5, 1, ["Acuerdo marco", "Catálogo", "Categoría", "Región / Departamento", "Provincia",
                   "Fichas", "Estado", "Detalle"], datos_det, col_estado=7, claves_estado=fills_det)
    ws5.freeze_panes = "A2"
    ws5.auto_filter.ref = ws5.dimensions
    auto_ancho(ws5)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()