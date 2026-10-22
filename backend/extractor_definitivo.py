import asyncio
from playwright.async_api import async_playwright
from bs4 import BeautifulSoup
import os
from openpyxl import Workbook, load_workbook

# Archivo final en formato Excel puro (.xlsx)
ARCHIVO_EXCEL = "convocatorias_seace.xlsx"

def guardar_en_excel_real(datos_fila):
    """
    Escribe directamente en un archivo binario de Excel (.xlsx) fila por fila.
    """
    try:
        if not os.path.exists(ARCHIVO_EXCEL):
            wb = Workbook()
            ws = wb.active
            ws.title = "Convocatorias SEACE"
            # Escribimos los encabezados principales
            ws.append([
                "N°", "Nombre o Sigla de Entidad", "Tipo de Selección", 
                "Nro. Selección", "Descripción del Objeto", "Año", 
                "Monto / Valor", "Estado / Fecha", "Datos Adicionales"
            ])
            wb.save(ARCHIVO_EXCEL)
        
        # Cargamos el archivo existente y añadimos la nueva convocatoria
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb.active
        ws.append(datos_fila)
        wb.save(ARCHIVO_EXCEL)
    except Exception as e:
        print(f"❌ Error al guardar físicamente en el archivo de Excel: {e}")

def limpiar_y_mostrar_tabla(xml_texto):
    """
    Función que limpia el XML del SEACE y extrae los datos de la tabla.
    """
    try:
        soup_xml = BeautifulSoup(xml_texto, "xml")
        actualizaciones = soup_xml.find_all("update")
        
        for update in actualizaciones:
            update_id = update.get("id", "")
            
            # FILTRO CORREGIDO: Escuchamos tanto la tabla principal como las grillas parciales de paginación (dtProcesos)
            if "pnlGrdResultadosProcesos" in update_id or "dtProcesos" in update_id:
                html_interno = update.text
                soup_html = BeautifulSoup(html_interno, "html.parser")
                
                # Buscamos las filas de datos reales
                filas = soup_html.find_all("tr")
                
                if len(filas) <= 1:
                    continue
                    
                print("\n========================================================")
                print(f"📊 [PROCESANDO NUEVOS DATOS] Se encontraron {len(filas)} filas en la tabla.")
                print("========================================================")
                
                contador_real = 1
                for idx, fila in enumerate(filas):
                    # Capturamos TODAS las columnas (celdas td) de la fila sin restricción
                    columnas = [td.get_text(strip=True) for td in fila.find_all("td")]
                    
                    # Filtramos que sean datos reales de convocatorias (más de 3 columnas de datos)
                    if columnas and len(columnas) > 3:
                        texto_linea = " | ".join(columnas)
                        
                        # Evitamos registrar las cabeceras repetidas o textos de error
                        if "Código Unico" in texto_linea or "No se encontraron" in texto_linea:
                            continue
                            
                        # 1. Mostramos el resultado completo en la consola limpia
                        print(f"🔹 Convocatoria #{contador_real}: {texto_linea[:120]}...")
                        
                        # 2. EXPORTACIÓN EXCEL REAL: Guarda la lista de datos nativa en el .xlsx
                        guardar_en_excel_real(columnas)
                        
                        contador_real += 1
                        
                print(f"💾 ¡Datos guardados y acumulados con éxito en el Excel '{ARCHIVO_EXCEL}'!")
                
    except Exception as e:
        print(f"❌ Error al procesar el texto: {e}")

async def interceptar_seace(response):
    """
    Esta función corre en segundo plano a velocidad extrema.
    Cada vez que el SEACE actualiza la tabla de procesos, de cualquier página, atrapamos los datos.
    """
    if "buscadorPublico.xhtml" in response.url:
        try:
            texto_respuesta = await response.text()
            
            # SOLUCIÓN DE PAGINACIÓN: Escuchamos tanto la actualización inicial como las parciales de PrimeFaces
            if "pnlGrdResultadosProcesos" in texto_respuesta or "dtProcesos" in texto_respuesta:
                print("\n📥 [¡DATOS CAPTURADOS!] Se detectó una actualización de la tabla de red.")
                limpiar_y_mostrar_tabla(texto_respuesta)
                
        except Exception:
            pass

async def main():
    async with async_playwright() as p:
        # Mantenemos tu configuración exacta que SÍ te abre la página
        browser = await p.chromium.launch(headless=False)
        context = await browser.new_context()
        page = await context.new_page()
        
        page.on("response", interceptar_seace)
        
        print("🔗 Navegando al buscador público del SEACE...")
        await page.goto("https://prod2.seace.gob.pe/seacebus-uiwd-pub/buscadorPublico/buscadorPublico.xhtml")
        
        print("\n💡 INSTRUCCIONES:")
        print("1. En la ventana del navegador que se abrió, llena los filtros de búsqueda.")
        print("2. Haz clic en el botón 'Buscar'.")
        print("3. Cuando veas la primera tabla en la pantalla, regresa aquí a la consola y presiona ENTER.")
        
        # PASO CLAVE INTERACTIVO: El script se detiene aquí hasta que tú le des ENTER en la terminal
        input("\n👉 Presiona la tecla [ENTER] aquí en la consola cuando ya veas los resultados en la web...")
        
        print("\n🚀 [PAGINACIÓN AUTOMÁTICA ACTIVADA] Avanzando de página solo...")
        
        pag_actual = 1
        while True:
            # Localizamos el botón 'Siguiente' activo del paginador de PrimeFaces (.ui-paginator-next)
            boton_siguiente = page.locator(".ui-paginator-next:not(.ui-state-disabled)").first
            
            # Si el botón ya no existe o está deshabilitado, rompemos el ciclo
            if await boton_siguiente.count() == 0:
                print("\n🏁 [FIN DEL PROCESO] Se llegó al final del catálogo o no hay más páginas.")
                break
                
            try:
                # 🛑 CAPTURA DE ANCLA: Tomamos el texto de la primera fila actual antes de dar clic
                primer_registro = page.locator("tr.ui-widget-content").first
                texto_anterior = ""
                if await primer_registro.count() > 0:
                    texto_anterior = await primer_registro.inner_text()

                print(f"🔄 Avanzando automáticamente a la siguiente página (Saliendo de Pág. {pag_actual})...")
                # El bot hace el clic de forma automática en la web
                await boton_siguiente.click()
                
                # ⏳ ESPERA DINÁMICA: Bloqueamos el avance hasta que la pantalla cambie de verdad su contenido
                cambio_exitoso = False
                for _ in range(20): # Espera máxima de 10 segundos por página (20 intentos x 0.5s)
                    await asyncio.sleep(0.5)
                    if await primer_registro.count() > 0:
                        texto_nuevo = await primer_registro.inner_text()
                        if texto_nuevo != texto_anterior:
                            cambio_exitoso = True
                            break
                
                # Un breve tiempo de colchón para que el interceptor termine de escribir de forma segura
                await page.wait_for_timeout(1500)
                pag_actual += 1
                
            except Exception as e:
                print(f"⚠️ Alerta en cambio de página: {e}. Esperando reintento...")
                await page.wait_for_timeout(3000)
                
        print(f"\n🎉 ¡Extracción finalizada! Revisa tu archivo Excel real '{ARCHIVO_EXCEL}'")

if __name__ == "__main__":
    asyncio.run(main())
